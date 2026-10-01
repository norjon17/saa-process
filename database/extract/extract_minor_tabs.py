"""Extract the small tabs -> seed/low_cpk.csv, circuit_burn.csv, fit_ocap.csv, weekly_ooc.csv

Each output has snake_case headers derived from the sheet's header row plus
source_row, tool_raw, tool_code (and `extra` for FIT OCAP, holding any column
whose header is not one of the documented ones).
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    Context, add_common_args, blank, cell_text, fmt_cell, iter_rows_until_gap, normalize_shift, normalize_tool,
    sheet_by_name, snake, to_iso_date, write_csv,
)

# sheet -> (header row, expected headers in order, output file, tool header, keep-row-if-any-of headers)
TABS = {
    "Low Cpk": dict(
        header_row=2, out="low_cpk.csv",
        expected=["No", "Date", "Shift", "SAA Tool", "Model/HT", "Lot Number", "Mean", "Sigma", "Cpk"],
        tool="SAA Tool", keep_if=["SAA Tool", "Lot Number"], extra=False, source="low_cpk"),
    "Circuit Burn": dict(
        header_row=2, out="circuit_burn.csv",
        expected=["No", "Date", "Shift", "SAA Tool", "Model/HT", "Lot Number", "Brick Qty", "Defect Qty", "Defect Rate"],
        tool="SAA Tool", keep_if=["SAA Tool", "Lot Number"], extra=False, source="circuit_burn"),
    "FIT OCAP": dict(
        header_row=1, out="fit_ocap.csv",
        expected=["No", "Date", "Shift", "SAA Tool", "Model", "FIT Detection", "Defect Name", "Brick Qty", "Defect Qty",
                  "Defect Rate", "FIT Confirmation", "Before", "After", "Remarks"],
        tool="SAA Tool", keep_if=None, extra=True, source="fit_ocap"),
    "Weekly OOC": dict(
        header_row=1, out="weekly_ooc.csv",
        expected=["No.", "WW", "Tool", "Model", "Head Type", "Suspension", "Issue", "Endorse by", "Endorse to", "Remarks After Action"],
        tool="Tool", keep_if=None, extra=False, source="weekly_ooc"),
}
DATE_HEADERS = {"date"}
SHIFT_HEADERS = {"shift"}


def extract_tab(wb, sheet: str, spec: dict, out_dir, ctx: Context) -> int:
    ws = sheet_by_name(wb, sheet)
    hdr_row = spec["header_row"]
    cols: list[tuple[int, str, str]] = []      # (index0, raw header, snake)
    extra_cols: list[tuple[int, str]] = []     # (index0, key) for FIT OCAP unnamed/undocumented columns
    expected_keys = {snake(h) for h in spec["expected"]}
    rows_out = []
    n_kept = 0
    keep_if_idx: list[int] = []
    tool_idx = None
    for n, r in iter_rows_until_gap(ws):
        if n < hdr_row:
            continue
        if n == hdr_row:
            names = [cell_text(v) for v in r]
            found = [h for h in names if h]
            missing = [h for h in spec["expected"] if h not in found]
            if missing:
                ctx.warn(f"{sheet}: expected header(s) {missing} not found; header row is {found}")
            for i, h in enumerate(names):
                if not h:
                    continue
                key = snake(h)
                if key in expected_keys or not spec["extra"]:
                    cols.append((i, h, key))
                else:
                    extra_cols.append((i, h))
            if spec["extra"]:
                # unnamed columns inside the header span also go to extra (keyed by column letter)
                last = max(i for i, _h, _k in cols) if cols else 0
                for i in range(last + 1):
                    if not names[i] and i not in {c[0] for c in cols} and i not in {c[0] for c in extra_cols}:
                        extra_cols.append((i, get_column_letter(i + 1)))
                unnamed_extra = [k for _i, k in extra_cols]
                if unnamed_extra:
                    ctx.decide(f"{sheet}: columns {unnamed_extra} are not in the documented header -> `extra` JSON")
            keys = [k for _i, _h, k in cols]
            if len(set(keys)) != len(keys):
                ctx.warn(f"{sheet}: duplicate snake_case header keys {keys}")
            tool_idx = next((i for i, h, _k in cols if h == spec["tool"]), None)
            if tool_idx is None:
                ctx.warn(f"{sheet}: tool column {spec['tool']!r} not found")
            if spec["keep_if"]:
                keep_if_idx = [i for i, h, _k in cols if h in spec["keep_if"]]
            else:
                # keep rows with any value outside the running-number column
                keep_if_idx = [i for i, h, _k in cols if snake(h) not in ("no",)] + [i for i, _k in extra_cols]
                ctx.decide(f"{sheet}: rows that only carry the running number (template filler) are not emitted")
            continue
        if not any(i < len(r) and not blank(r[i]) and fmt_cell(r[i]) != "" for i in keep_if_idx):
            continue
        where = f"{sheet}!row {n}"
        tool_raw = r[tool_idx] if tool_idx is not None and tool_idx < len(r) else None
        tool_code = ""
        if not blank(tool_raw):
            tool_code = normalize_tool(tool_raw, ctx) or ""
            if not tool_code:
                ctx.warn(f"{where}: unresolved tool {tool_raw!r}")
            else:
                ctx.see_tool(spec["source"], tool_code, raw=cell_text(tool_raw))
        rec = [n, cell_text(tool_raw), tool_code]
        for i, h, key in cols:
            v = r[i] if i < len(r) else None
            if key in DATE_HEADERS:
                rec.append(to_iso_date(v, ctx, where))
            elif key in SHIFT_HEADERS:
                rec.append(normalize_shift(v, ctx, where))
            else:
                rec.append(fmt_cell(v))
        if spec["extra"]:
            extra = {k: fmt_cell(r[i]) for i, k in extra_cols if i < len(r) and not blank(r[i]) and fmt_cell(r[i]) != ""}
            rec.append(json.dumps(extra, ensure_ascii=False) if extra else "")
        rows_out.append(rec)
        n_kept += 1
    header = ["source_row", "tool_raw", "tool_code"] + [k for _i, _h, k in cols] + (["extra"] if spec["extra"] else [])
    out = Path(out_dir) / spec["out"]
    n = write_csv(out, header, rows_out)
    ctx.count(spec["out"], n)
    print(f"  {spec['out']}: {n} rows (from sheet {sheet!r})")
    return n


def run(src_dir, out_dir, ctx: Context) -> dict[str, int]:
    wb = ctx.endorsement_wb(src_dir)
    return {sheet: extract_tab(wb, sheet, spec, out_dir, ctx) for sheet, spec in TABS.items()}


def main(argv=None):
    ap = add_common_args(argparse.ArgumentParser(description=__doc__))
    a = ap.parse_args(argv)
    ctx = Context()
    run(a.src, a.out, ctx)
    ctx.save("minor_tabs")
    print("  unmapped:", {k: dict(v) for k, v in ctx.unmapped.items() if v})


if __name__ == "__main__":
    main()
