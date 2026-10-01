"""Extract every 'PM STATUS <MONTH> <YEAR>' sheet -> seed/pm_status.csv

Columns are located by header text (rows 1-3, merged group headers forward-filled);
the resolved layout is compared with the expected fixed layout and any difference
is reported as a warning.  Unknown columns with data go into the `extra` JSON column.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

from openpyxl.utils import column_index_from_string, get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    Context, add_common_args, blank, cell_text, fmt_cell, fmt_flag, fmt_num, iter_rows_until_gap,
    normalize_region, normalize_tool, split_people, to_iso_date, write_csv,
)

SHEET_RE = re.compile(r"^\s*PM STATUS\s+([A-Z]+)\s+(\d{4})\s*$", re.IGNORECASE)
MONTHS = {m.casefold(): i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"], 1)}

HEADER = [
    "sheet", "period", "tool_raw", "tool_code", "area", "tool_type", "ee_raw", "ee_id_no", "actual_date",
    "fix_vac_l_with", "fix_vac_l_without", "fix_vac_r_with", "fix_vac_r_without",
    "pick_vac_l_with", "pick_vac_l_without", "pick_vac_r_with", "pick_vac_r_without",
    "clamp_l_clamp", "clamp_l_unclamp", "clamp_r_clamp", "clamp_r_unclamp",
    "fixture_spring", "speed_motor",
    "z_off_l_pickup", "z_off_l_release", "z_off_r_pickup", "z_off_r_release",
    "opt_grr", "opt_correlation", "opt_laser_recipe", "opt_additional_area", "opt_scan_pattern_dir",
    "fixture_left", "fixture_right", "scan_pattern", "extra",
]

# field -> (expected column letter, leaf header text, required group texts, formatter)
NUM, FLAG, TXT, DATE = "num", "flag", "txt", "date"
FIELDS = {
    "area":               ("A",  "area", (), TXT),
    "tool_raw":           ("B",  "tool number", (), TXT),
    "tool_type":          ("C",  "tool type", (), TXT),
    "ee_raw":             ("D",  "ee incharge", (), TXT),
    "actual_date":        ("E",  "actual date", (), DATE),
    "fix_vac_l_with":     ("F",  "with hga", ("left fixture vacuum",), NUM),
    "fix_vac_l_without":  ("G",  "w/o hga", ("left fixture vacuum",), NUM),
    "fix_vac_r_with":     ("J",  "with hga", ("right fixture vacuum",), NUM),
    "fix_vac_r_without":  ("K",  "w/o hga", ("right fixture vacuum",), NUM),
    "pick_vac_l_with":    ("N",  "with hga", ("left picker vacuum",), NUM),
    "pick_vac_l_without": ("O",  "w/o hga", ("left picker vacuum",), NUM),
    "pick_vac_r_with":    ("R",  "with hga", ("right picker vacuum",), NUM),
    "pick_vac_r_without": ("S",  "w/o hga", ("right picker vacuum",), NUM),
    "clamp_l_clamp":      ("V",  "clamp", ("fixture clamping speed", "left"), NUM),
    "clamp_l_unclamp":    ("W",  "unclamp", ("fixture clamping speed", "left"), NUM),
    "clamp_r_clamp":      ("X",  "clamp", ("fixture clamping speed", "right"), NUM),
    "clamp_r_unclamp":    ("Y",  "unclamp", ("fixture clamping speed", "right"), NUM),
    "fixture_spring":     ("Z",  "fixture spring", (), NUM),
    "speed_motor":        ("AA", "speed motor", (), NUM),
    "z_off_l_pickup":     ("AB", "pick up", ("z height offset (fixture)", "left"), NUM),
    "z_off_l_release":    ("AC", "release", ("z height offset (fixture)", "left"), NUM),
    "z_off_r_pickup":     ("AD", "pick up", ("z height offset (fixture)", "right"), NUM),
    "z_off_r_release":    ("AE", "release", ("z height offset (fixture)", "right"), NUM),
    "opt_grr":            ("AF", "grr", ("optimization items",), FLAG),
    "opt_correlation":    ("AG", "correlation", ("optimization items",), FLAG),
    "opt_laser_recipe":   ("AH", "laser recipe", ("optimization items",), FLAG),
    "opt_additional_area": ("AI", "additional area/lines", ("optimization items",), FLAG),
    "opt_scan_pattern_dir": ("AJ", "scan pattern direction", ("optimization items",), FLAG),
    "fixture_left":       ("AK", "fixture left", (), TXT),
    "fixture_right":      ("AL", "fixture right", (), TXT),
    "scan_pattern":       ("AM", "scan pattern", (), TXT),
}
IGNORED_LEAVES = {"h1", "delta"}  # formula columns (avg / delta) are not extracted


def _norm(s) -> str:
    return re.sub(r"\s+", " ", cell_text(s)).casefold()


def resolve_layout(hdr_rows: list[tuple], ctx: Context, sheet: str):
    """Return (field->col_index0, extra_cols: {col_index0: key}, ignored: set)."""
    ncol = max(len(r) for r in hdr_rows)
    raw = [[_norm(r[i]) if i < len(r) else "" for i in range(ncol)] for r in hdr_rows]
    r1, r2, r3 = raw
    f1, f2 = [""] * ncol, [""] * ncol
    cur1 = cur2 = ""
    for i in range(ncol):
        if r1[i]:
            cur1, cur2 = r1[i], ""
        if r2[i]:
            cur2 = r2[i]
        f1[i], f2[i] = cur1, cur2
    leaf = [r3[i] or r2[i] or r1[i] for i in range(ncol)]

    mapping: dict[str, int] = {}
    used: set[int] = set()
    for field, (exp_letter, leaf_txt, groups, _fmt) in FIELDS.items():
        cands = [i for i in range(ncol) if leaf[i] == leaf_txt and all(g in (f1[i], f2[i]) for g in groups) and i not in used]
        if not cands:
            ctx.warn(f"{sheet}: header for {field} ({leaf_txt!r} in {groups}) not found -> column blank")
            continue
        i = cands[0]
        if len(cands) > 1:
            ctx.warn(f"{sheet}: header for {field} matched several columns {[get_column_letter(c + 1) for c in cands]}; using first")
        mapping[field] = i
        used.add(i)
        if get_column_letter(i + 1) != exp_letter:
            ctx.warn(f"{sheet}: layout differs - {field} found in column {get_column_letter(i + 1)} (expected {exp_letter})")
    ignored = {i for i in range(ncol) if i not in used and leaf[i] in IGNORED_LEAVES}
    extra = {}
    for i in range(ncol):
        if i in used or i in ignored:
            continue
        parts = [p for p in (r1[i], r2[i], r3[i]) if p]
        extra[i] = " / ".join(parts) if parts else get_column_letter(i + 1)
    return mapping, extra, ignored


def period_of(title: str) -> str | None:
    m = SHEET_RE.match(title)
    if not m:
        return None
    mon = MONTHS.get(m.group(1).casefold())
    if mon is None:
        return None
    return dt.date(int(m.group(2)), mon, 1).isoformat()


def run(src_dir, out_dir, ctx: Context) -> int:
    wb = ctx.endorsement_wb(src_dir)
    rows_out = []
    sheets_done = []
    for ws in wb.worksheets:
        title = ws.title
        if not re.match(r"^\s*PM STATUS\b", title, re.IGNORECASE):
            if re.match(r"^\s*(JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\s+\d{4}\s*$", title, re.I):
                ctx.decide(f"PM STATUS: sheet {title!r} skipped (not a 'PM STATUS <MONTH> <YEAR>' tab; it is empty)")
            continue
        period = period_of(title)
        if not period:
            ctx.warn(f"PM STATUS sheet {title!r}: cannot parse month/year from title -> skipped")
            continue
        hdr_rows: list[tuple] = []
        per_sheet: dict[str, list] = {}   # tool_code -> row (dedup)
        per_sheet_filled: dict[str, int] = {}
        n_rows = 0
        mapping = extra_cols = ignored = None
        for n, r in iter_rows_until_gap(ws):
            if n <= 3:
                hdr_rows.append(r)
                if n == 3:
                    mapping, extra_cols, ignored = resolve_layout(hdr_rows, ctx, title)
                continue
            if mapping is None:
                ctx.warn(f"{title}: data before header row 3 - skipped")
                break
            get = lambda f: (r[mapping[f]] if f in mapping and mapping[f] < len(r) else None)  # noqa: E731
            tool_raw = get("tool_raw")
            if blank(tool_raw):
                # a row without a tool: skip; report if it has other content
                if any(not blank(v) for v in r):
                    ctx.warn(f"{title}!row {n}: row without tool number skipped")
                continue
            where = f"{title}!row {n}"
            tool_code = normalize_tool(tool_raw, ctx)
            if tool_code is None:
                ctx.warn(f"{where}: unresolved tool {tool_raw!r} -> row skipped")
                continue
            n_rows += 1
            area = normalize_region(get("area"), ctx, where)
            ee_raw = cell_text(get("ee_raw"))
            ee_ids = split_people(ee_raw, ctx) if ee_raw else []
            rec = {
                "sheet": title, "period": period, "tool_raw": cell_text(tool_raw), "tool_code": tool_code,
                "area": area, "tool_type": cell_text(get("tool_type")).upper(), "ee_raw": ee_raw,
                "ee_id_no": "|".join(ee_ids), "actual_date": to_iso_date(get("actual_date"), ctx, where),
            }
            for field, (_l, _t, _g, fmt) in FIELDS.items():
                if field in rec:
                    continue
                v = get(field)
                rec[field] = fmt_num(v) if fmt == NUM else fmt_flag(v) if fmt == FLAG else fmt_cell(v)
            extra = {}
            for i, key in extra_cols.items():
                if i < len(r) and not blank(r[i]):
                    extra[key] = fmt_cell(r[i]) or cell_text(r[i])
            rec["extra"] = json.dumps(extra, ensure_ascii=False) if extra else ""
            row = [rec[h] for h in HEADER]
            filled = sum(1 for h in HEADER if rec[h] not in ("", None))
            if tool_code in per_sheet:
                prev_filled = per_sheet_filled[tool_code]
                keep_new = filled > prev_filled
                ctx.warn(f"{title}: tool {tool_code} appears twice (raw {per_sheet[tool_code][2]!r} / {rec['tool_raw']!r}); "
                         f"kept the row with more filled cells ({'new' if keep_new else 'first'}: {max(filled, prev_filled)} vs {min(filled, prev_filled)})")
                if not keep_new:
                    continue
            per_sheet[tool_code] = row
            per_sheet_filled[tool_code] = filled
            ctx.see_tool("pm_status", tool_code, raw=cell_text(tool_raw), region=area or None,
                         tool_type=rec["tool_type"] or None, ee=None, period=None)
            # record EE per (tool, period) for extract_master ("EE from the latest month that has one")
            info = ctx.tools["pm_status"][tool_code]
            info.setdefault("by_period", {})[period] = {"ee": ee_raw, "ee_ids": ee_ids, "region": area, "tool_type": rec["tool_type"]}
        rows_out.extend(per_sheet.values())
        sheets_done.append((title, period, len(per_sheet)))
        print(f"  {title}: {len(per_sheet)} tools (period {period})")
    rows_out.sort(key=lambda r: (r[1], r[3]))
    out = Path(out_dir) / "pm_status.csv"
    n = write_csv(out, HEADER, rows_out)
    ctx.count("pm_status.csv", n)
    ctx.count("pm_status.csv sheets", len(sheets_done))
    print(f"  pm_status.csv: {n} rows from {len(sheets_done)} sheets")
    return n


def main(argv=None):
    ap = add_common_args(argparse.ArgumentParser(description=__doc__))
    a = ap.parse_args(argv)
    ctx = Context()
    run(a.src, a.out, ctx)
    ctx.save("pm_status")
    print("  unmapped:", {k: dict(v) for k, v in ctx.unmapped.items() if v})


if __name__ == "__main__":
    main()
