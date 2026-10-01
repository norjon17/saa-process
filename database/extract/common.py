"""Shared helpers for the one-time seed extraction.

Everything here is pure Python + openpyxl.  The source workbooks are opened
READ-ONLY and never written.

Public surface used by the extract_*.py scripts:
    Context            accumulates unmapped values, warnings and the per-source tool sets
    normalize_tool     "SAA 54H" / "SAA05" / "SA53H" -> "SAA54H" / "SAA05H" / "SAA53H"
    split_people       "Coleen/Sam " -> ["9006", "9005"]
    normalize_category, normalize_shift, normalize_progress, normalize_region
    to_iso_date, fmt_cell
    write_csv, write_json, read_json
    open_endorsement_workbook, iter_rows_until_gap
"""
from __future__ import annotations

import csv
import datetime as _dt
import json
import os
import re
from collections import Counter, OrderedDict
from pathlib import Path

# --------------------------------------------------------------------------- paths
EXTRACT_DIR = Path(__file__).resolve().parent
DATABASE_DIR = EXTRACT_DIR.parent
DEFAULT_SEED_DIR = DATABASE_DIR / "seed"
DEFAULT_SRC_DIR = Path(r"D:\Workstation\evalice\docs")
INTERMEDIATE_DIR = EXTRACT_DIR / ".intermediate"

ENDORSEMENT_XLSX = Path("Endorsement nmin n nkaexcel") / "SAA Process Endorsement_2026_NEW.xlsx"
MEANCHECK_XLSX = Path("all") / "all about" / "1st Meancheck excel template" / "First Mean Check Monitoring Rev2.1_NEW.xlsx"


def src_path(src_dir: str | os.PathLike, rel: Path) -> Path:
    p = Path(src_dir) / rel
    if not p.exists():
        raise FileNotFoundError(f"source workbook not found: {p}")
    return p


# --------------------------------------------------------------------------- master lists
# The master user list.  Written verbatim to seed/users.json by extract_master.py
# and used here to build the alias -> idNo map.
USERS = [
    {"idNo": "1111", "name": "Admin", "role": "admin", "aliases": []},
    {"idNo": "9000", "name": "Evalice Villacampa", "role": "admin", "aliases": ["evalice"]},
    {"idNo": "9001", "name": "Mary Rose Alis", "role": "admin", "aliases": ["mary rose", "rose"]},
    {"idNo": "9002", "name": "Eva", "role": "tech", "aliases": ["eva"]},
    {"idNo": "9003", "name": "Dea", "role": "tech", "aliases": ["dea"]},
    {"idNo": "9004", "name": "KC", "role": "tech", "aliases": ["kc", "k.c"]},
    {"idNo": "9005", "name": "Sam", "role": "tech", "aliases": ["sam"]},
    {"idNo": "9006", "name": "Coleen", "role": "tech", "aliases": ["coleen", "collen"]},
    {"idNo": "9007", "name": "Camille", "role": "tech", "aliases": ["camille", "camiile"]},
    {"idNo": "9008", "name": "Fil", "role": "tech", "aliases": ["fil"]},
    {"idNo": "9009", "name": "Vinz", "role": "tech", "aliases": ["vinz"]},
    {"idNo": "9010", "name": "Rona", "role": "tech", "aliases": ["rona"]},
    {"idNo": "9011", "name": "Ralf Anthony David", "role": "ee", "aliases": ["ralf anthony david", "ralf", "raf"]},
    {"idNo": "9012", "name": "John Joseph Mingo", "role": "ee", "aliases": ["john joseph mingo", "jj"]},
    {"idNo": "9013", "name": "Hileton Mago", "role": "ee", "aliases": ["hileton mago", "hileton"]},
    {"idNo": "9014", "name": "Michael Monterey Javier", "role": "ee", "aliases": ["michael monterey javier", "michael"]},
    {"idNo": "9015", "name": "Ramonchito Juanillas", "role": "ee", "aliases": ["ramonchito juanillas", "ramonchito", "chito"]},
    {"idNo": "9016", "name": "Glenn Rubis Jr.", "role": "ee", "aliases": ["glenn rubis jr.", "glenn rubis jr", "glenn rubis", "glenn"]},
    {"idNo": "9017", "name": "Jay Bola\u00f1os", "role": "ee", "aliases": ["jay bola\u00f1os", "jay bolanos", "jay bola\ufffdos", "jay"]},
    {"idNo": "9018", "name": "Arnel Galve", "role": "ee", "aliases": ["arnel galve", "arnel"]},
    {"idNo": "9019", "name": "Joshua Natanauan", "role": "ee", "aliases": ["joshua natanauan", "joshua"]},
    {"idNo": "9020", "name": "Vince", "role": "ee", "aliases": ["vince"]},
    {"idNo": "9021", "name": "Noel", "role": "ee", "aliases": ["noel"]},
    {"idNo": "9022", "name": "Antonio (Ton)", "role": "ee", "aliases": ["antonio", "ton"]},
]

REGIONS = ["Luzon", "Mindanao", "Visayas", "CB Line", "NPI", "Class 100"]

CATEGORIES = [
    ("Low Yield", "s1"),
    ("High Offset", "s1"),
    ("Parametric Issue", "s1"),
    ("Mechanical Issue", "s2"),
    ("Conversion", "s3"),
    ("New Product", "s3"),
    ("Optimization", "s4"),
    ("PM", "s5"),
    ("Mean Check Confirmation", "s6"),
    ("Static SPC OOC", "s6"),
]

# code, pitch_target, roll_target, tolerance  (limits are the same for every model)
MODELS = [
    ("J1B", 1.8, 0, 0.4), ("J1R", 1.8, 0, 0.4), ("JGO", 1.8, 0, 0.4), ("JGB", 1.8, 0, 0.4),
    ("MPK", 2.0, 0, 0.4), ("SF4", 1.9, 0, 0.4), ("CEP", 1.7, 0, 0.65), ("A8S", 2.0, 0, 0.4),
    ("APS", 2.0, 0, 0.4), ("PCM", 1.9, 0, 0.4), ("PDQ", 1.9, 0, 0.4), ("LDS", 1.9, 0, 0.4),
    ("SEW", 1.9, 0, 0.4), ("LD3", 1.9, 0, 0.4), ("KCF", 1.75, 0, 0.25),
]
MODEL_LIMITS = {"delta_sat_limit": 0.05, "delta_target_limit": 0.05, "sigma_limit": 0.08, "cpk_min": 1.33}

# Raw tool spellings that the SAA regex alone cannot resolve.
TOOL_ALIASES = {
    "SA53H": "SAA53H",
    "SA72H": "SAA72H",
    "SA79H": "SAA79H",
    "SASA33": "SAA33H",
    "SAA677H": "SAA77H",
}

# Area / region spellings seen across the workbooks -> canonical region names.
REGION_ALIASES = {
    "luzon": "Luzon",
    "mindanao": "Mindanao",
    "visayas": "Visayas",
    "cb line": "CB Line",
    "cbline": "CB Line",
    "cb13": "CB Line",
    "cb14": "CB Line",
    "npi": "NPI",
    "class 100": "Class 100",
    "class100": "Class 100",
}

DAILY_CELL_VALUES = {
    "p (auto)": "P_AUTO",
    "f (auto)": "F_AUTO",
    "p (man)": "P_MAN",
    "f (man)": "F_MAN",
    "idle": "IDLE",
    "mb": "MB",
}


# --------------------------------------------------------------------------- context
class Context:
    """Accumulates unmapped values, warnings and per-source tool info across extractors."""

    UNMAPPED_KEYS = ("tools", "people", "categories", "contact_chip_values", "daily_cells")

    def __init__(self) -> None:
        self.unmapped: dict[str, Counter] = {k: Counter() for k in self.UNMAPPED_KEYS}
        self.warnings: list[str] = []
        self.decisions: list[str] = []
        # tools[source][code] = {"raw": [...], "region": .., "tool_type": .., "ee": .., "period": ..}
        self.tools: dict[str, dict[str, dict]] = {}
        # raw spelling (upper, whitespace removed) -> code, for spellings resolved via TOOL_ALIASES
        self.tool_alias_hits: dict[str, str] = {}
        self.counts: OrderedDict[str, int] = OrderedDict()
        self._wb_cache: dict[str, object] = {}

    # -- reporting
    def warn(self, msg: str) -> None:
        self.warnings.append(msg)
        print("  WARN:", msg)

    def decide(self, msg: str) -> None:
        self.decisions.append(msg)

    def count(self, name: str, n: int) -> None:
        self.counts[name] = n

    # -- tools
    def see_tool(self, source: str, code: str, raw=None, region=None, tool_type=None, ee=None, period=None) -> None:
        d = self.tools.setdefault(source, {}).setdefault(code, {"raw": []})
        if raw is not None and str(raw) not in d["raw"]:
            d["raw"].append(str(raw))
        if region:
            d["region"] = region
        if tool_type:
            d["tool_type"] = tool_type
        if ee:
            d["ee"] = ee
        if period:
            d["period"] = period

    # -- persistence (so extractors can run standalone and extract_master can pick up their tool sets)
    def to_dict(self) -> dict:
        return {
            "unmapped": {k: dict(v) for k, v in self.unmapped.items()},
            "warnings": list(self.warnings),
            "decisions": list(self.decisions),
            "tools": self.tools,
            "tool_alias_hits": self.tool_alias_hits,
            "counts": dict(self.counts),
        }

    def save(self, name: str) -> Path:
        INTERMEDIATE_DIR.mkdir(parents=True, exist_ok=True)
        p = INTERMEDIATE_DIR / f"{name}.json"
        write_json(p, self.to_dict())
        return p

    def merge_dict(self, d: dict) -> None:
        for k, v in d.get("unmapped", {}).items():
            self.unmapped.setdefault(k, Counter()).update(v)
        for w in d.get("warnings", []):
            if w not in self.warnings:
                self.warnings.append(w)
        for w in d.get("decisions", []):
            if w not in self.decisions:
                self.decisions.append(w)
        for src, tools in d.get("tools", {}).items():
            for code, info in tools.items():
                cur = self.tools.setdefault(src, {}).setdefault(code, {"raw": []})
                for r in info.get("raw", []):
                    if r not in cur["raw"]:
                        cur["raw"].append(r)
                for key in ("region", "tool_type", "ee", "period"):
                    if info.get(key):
                        cur[key] = info[key]
                if info.get("by_period"):
                    cur.setdefault("by_period", {}).update(info["by_period"])
        self.tool_alias_hits.update(d.get("tool_alias_hits", {}))
        self.counts.update(d.get("counts", {}))

    @classmethod
    def load_all(cls, names: list[str]) -> "Context":
        ctx = cls()
        for n in names:
            p = INTERMEDIATE_DIR / f"{n}.json"
            if p.exists():
                ctx.merge_dict(read_json(p))
            else:
                ctx.warn(f"intermediate file missing: {p} (run extract_{n}.py first)")
        return ctx

    # -- workbook cache
    def endorsement_wb(self, src_dir):
        key = str(src_path(src_dir, ENDORSEMENT_XLSX))
        if key not in self._wb_cache:
            self._wb_cache[key] = open_endorsement_workbook(key)
        return self._wb_cache[key]


# --------------------------------------------------------------------------- normalisers
_TOOL_RE = re.compile(r"^SAA(\d{1,2})H?$")
_TOOL_MISSING_A_RE = re.compile(r"^SA(\d{2})H$")


def _clean_ws(s) -> str:
    return re.sub(r"\s+", "", str(s))


def normalize_tool(raw, ctx: Context | None = None) -> str | None:
    """Return canonical SAAnnH code or None (and record in ctx.unmapped['tools'])."""
    if raw is None:
        return None
    s = _clean_ws(raw).upper()
    if not s:
        return None
    m = _TOOL_RE.match(s)
    if m:
        return f"SAA{int(m.group(1)):02d}H"
    # secondary resolutions: explicit alias table, stray trailing punctuation ("SAA45H-"),
    # a missing 'A' ("SA66H").  Every spelling resolved this way is recorded so that
    # tool_aliases.json lists it.
    code = TOOL_ALIASES.get(s)
    if code is None:
        t = s.rstrip("-.:;,_")
        m = _TOOL_RE.match(t) or _TOOL_MISSING_A_RE.match(t)
        if m:
            code = f"SAA{int(m.group(1)):02d}H"
        elif t in TOOL_ALIASES:
            code = TOOL_ALIASES[t]
    if code:
        if ctx is not None:
            ctx.tool_alias_hits[s] = code
        return code
    if ctx is not None:
        ctx.unmapped["tools"][str(raw)] += 1
    return None


def tool_sort_order(code: str) -> int:
    m = re.search(r"(\d+)", code or "")
    return int(m.group(1)) if m else 0


_ALIASES: dict[str, str] | None = None
_USERS_BY_ID: dict[str, dict] | None = None


def person_key(token) -> str:
    """casefold + strip, strip trailing dots, collapse internal whitespace."""
    s = str(token).strip().casefold()
    s = s.rstrip(".").strip()
    s = re.sub(r"\s+", " ", s)
    return s


def user_aliases() -> dict[str, str]:
    global _ALIASES, _USERS_BY_ID
    if _ALIASES is None:
        _ALIASES = {}
        _USERS_BY_ID = {}
        for u in USERS:
            _USERS_BY_ID[u["idNo"]] = u
            for a in u["aliases"]:
                _ALIASES[person_key(a)] = u["idNo"]
            _ALIASES.setdefault(person_key(u["name"]), u["idNo"])
    return _ALIASES


def user_name(id_no: str) -> str:
    user_aliases()
    return _USERS_BY_ID[id_no]["name"]


_PEOPLE_SEP = re.compile(r"\s*(?:/|\\|,|&|\+|\band\b)\s*", re.IGNORECASE)


def split_people(raw, ctx: Context | None = None) -> list[str]:
    """Split a free-text name cell into canonical user idNo values (order kept, de-duplicated)."""
    if raw is None:
        return []
    s = str(raw).strip()
    if not s:
        return []
    aliases = user_aliases()
    out: list[str] = []
    for tok in _PEOPLE_SEP.split(s):
        key = person_key(tok)
        if not key:
            continue
        idn = aliases.get(key)
        if idn is None:
            if ctx is not None:
                ctx.unmapped["people"][tok.strip()] += 1
            continue
        if idn not in out:
            out.append(idn)
    return out


_CATEGORY_KEYS = {name.casefold(): name for name, _ in CATEGORIES}


def normalize_category(raw, ctx: Context | None = None) -> str:
    if raw is None:
        return ""
    s = str(raw).strip()
    if not s:
        return ""
    name = _CATEGORY_KEYS.get(re.sub(r"\s+", " ", s).casefold())
    if name is None:
        if ctx is not None:
            ctx.unmapped["categories"][s] += 1
        return ""
    return name


def normalize_shift(raw, ctx: Context | None = None, where: str = "") -> str:
    if raw is None:
        return ""
    s = str(raw).strip().upper()
    if s in ("DS", "NS"):
        return s
    if s and ctx is not None:
        ctx.warn(f"unknown shift {raw!r} at {where} -> blank")
    return ""


def normalize_progress(raw) -> str:
    s = "" if raw is None else str(raw).strip().casefold()
    if s == "done":
        return "done"
    if s == "ongoing":
        return "ongoing"
    return "open"


def normalize_region(raw, ctx: Context | None = None, where: str = "") -> str:
    if raw is None:
        return ""
    s = re.sub(r"\s+", " ", str(raw).strip()).casefold()
    if not s:
        return ""
    r = REGION_ALIASES.get(s)
    if r is None:
        if s.startswith("class"):
            r = "Class 100"
        if ctx is not None:
            ctx.warn(f"unknown area/region {raw!r} at {where} -> {r or 'blank'}")
    return r or ""


def normalize_daily_cell(raw, ctx: Context | None = None) -> str:
    if raw is None:
        return ""
    s = str(raw).strip()
    if not s:
        return ""
    v = DAILY_CELL_VALUES.get(re.sub(r"\s+", " ", s).casefold())
    if v is None:
        if ctx is not None:
            ctx.unmapped["daily_cells"][s] += 1
        return ""
    return v


# --------------------------------------------------------------------------- value formatting
def to_iso_date(v, ctx: Context | None = None, where: str = "") -> str:
    """datetime/date -> YYYY-MM-DD.  Strings are parsed if they look like a date; else blank + warning."""
    if v is None:
        return ""
    if isinstance(v, _dt.datetime):
        return v.date().isoformat()
    if isinstance(v, _dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        from openpyxl.utils.datetime import from_excel
        try:
            return from_excel(v).date().isoformat()
        except Exception:
            pass
    s = str(v).strip()
    if not s:
        return ""
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%b %d, %Y", "%B %d, %Y"):
        try:
            return _dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    repaired = _repair_date_typo(s)
    if repaired:
        if ctx is not None:
            ctx.warn(f"typo date {v!r} at {where} -> repaired to {repaired}")
        return repaired
    if ctx is not None:
        ctx.warn(f"unparseable date {v!r} at {where} -> blank")
    return ""


def _repair_date_typo(s: str) -> str:
    """Salvage obvious M/D/YYYY typing errors (the workbooks use US order for real dates):
    '4/14//2025' (doubled slash), '5/212026' / '11/122025' (missing second slash).
    Anything ambiguous (e.g. a 5-digit year) is left to the caller to blank."""
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/+(\d{4})", s)
    if not m:
        m = re.fullmatch(r"(\d{1,2})/(\d{2})(\d{4})", s)
    if not m:
        return ""
    try:
        return _dt.date(int(m.group(3)), int(m.group(1)), int(m.group(2))).isoformat()
    except ValueError:
        return ""


def is_error(v) -> bool:
    return isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A", "L"))


def fmt_cell(v) -> str:
    """Excel cell -> CSV text.  Errors (#DIV/0! ...) become blank; floats are trimmed."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:.10g}"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, _dt.datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ")
    if isinstance(v, _dt.date):
        return v.isoformat()
    if isinstance(v, _dt.time):
        return v.isoformat()
    s = str(v)
    if is_error(s.strip()):
        return ""
    return s.strip()


def fmt_num(v) -> str:
    """Numeric cell -> text; anything non-numeric (incl. Excel errors) -> blank."""
    if v is None or isinstance(v, bool):
        return ""
    if isinstance(v, (int, float)):
        return fmt_cell(v)
    s = str(v).strip()
    try:
        return fmt_cell(float(s))
    except ValueError:
        return ""


def fmt_flag(v) -> str:
    """'O' -> true, 'X' -> false, blank -> ''."""
    if v is None:
        return ""
    s = str(v).strip().upper()
    if s == "O":
        return "true"
    if s == "X":
        return "false"
    return ""


def cell_text(v) -> str:
    return "" if v is None else str(v).strip()


def blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def snake(s: str) -> str:
    s = re.sub(r"[^0-9a-zA-Z]+", "_", str(s).strip()).strip("_").lower()
    return s or "col"


# --------------------------------------------------------------------------- IO
def write_csv(path, header: list[str], rows) -> int:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    n = 0
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_MINIMAL, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow(["" if x is None else x for x in r])
            n += 1
    return n


def write_json(path, data) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.write("\n")


def read_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def open_endorsement_workbook(path):
    import openpyxl
    return openpyxl.load_workbook(str(path), read_only=True, data_only=True)


def iter_rows_until_gap(ws, gap: int = 50, min_len: int = 0):
    """Yield (row_number, tuple_of_values) for a read-only sheet, stopping after `gap`
    consecutive completely-empty rows.  Never touches ws.max_row / ws.dimensions
    (the hidden 'Weekly OOC' sheet has ~1M formatted rows)."""
    empty = 0
    n = 0
    for r in ws.iter_rows(values_only=True):
        n += 1
        if all(blank(v) for v in r):
            empty += 1
            if empty >= gap:
                break
            continue
        empty = 0
        if min_len and len(r) < min_len:
            r = tuple(r) + (None,) * (min_len - len(r))
        yield n, r


def sheet_by_name(wb, name: str):
    if name in wb.sheetnames:
        return wb[name]
    # tolerate trailing/leading whitespace differences in sheet titles
    for n in wb.sheetnames:
        if n.strip().casefold() == name.strip().casefold():
            return wb[n]
    raise KeyError(f"sheet {name!r} not found; have {wb.sheetnames}")


def add_common_args(parser):
    parser.add_argument("--src", default=str(DEFAULT_SRC_DIR), help="docs dir containing the source workbooks")
    parser.add_argument("--out", default=str(DEFAULT_SEED_DIR), help="seed output dir")
    return parser
