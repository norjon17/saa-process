"""Extract the 'Contact Chip' sheet -> seed/contact_chip.csv

Row 1 holds dates from column D onward; columns A-C are region / tool / tool type.
Cell values L, R, L/R, PM, Repair -> L, R, LR, PM, REPAIR.  Any other text is kept
in `note` and reported in unmapped.contact_chip_values.
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    Context, add_common_args, blank, cell_text, iter_rows_until_gap, normalize_region, normalize_tool,
    sheet_by_name, to_iso_date, write_csv,
)

SHEET = "Contact Chip"
HEADER = ["tool_code", "date", "value", "note"]
VALUE_MAP = {"l": "L", "r": "R", "l/r": "LR", "lr": "LR", "r/l": "LR", "pm": "PM", "repair": "REPAIR"}


def map_value(raw, ctx: Context):
    s = re.sub(r"\s+", "", cell_text(raw)).casefold()
    v = VALUE_MAP.get(s)
    if v is None:
        ctx.unmapped["contact_chip_values"][cell_text(raw)] += 1
        return "", cell_text(raw)
    return v, ""


def run(src_dir, out_dir, ctx: Context) -> int:
    wb = ctx.endorsement_wb(src_dir)
    ws = sheet_by_name(wb, SHEET)
    dates: dict[int, str] = {}   # column index0 -> iso date
    rows = []
    seen_tools: dict[str, int] = {}
    values: Counter = Counter()
    for n, r in iter_rows_until_gap(ws):
        if n == 1:
            for i in range(3, len(r)):
                if blank(r[i]):
                    continue
                d = to_iso_date(r[i], ctx, f"{SHEET}!{get_column_letter(i + 1)}1")
                if d:
                    dates[i] = d
                else:
                    ctx.warn(f"{SHEET}!{get_column_letter(i + 1)}1: header cell {r[i]!r} is not a date -> column ignored")
            if not dates:
                ctx.warn(f"{SHEET}: no dates found in row 1")
            continue
        tool_raw = r[1] if len(r) > 1 else None
        if blank(tool_raw):
            continue
        # legend rows ('Repair', 'PM', 'Left Fixture' ...) do not normalise to a tool code: skip silently
        s = re.sub(r"\s+", "", str(tool_raw)).upper()
        if not re.match(r"^SA", s):
            continue
        tool_code = normalize_tool(tool_raw, ctx)
        if tool_code is None:
            ctx.warn(f"{SHEET}!row {n}: unresolved tool {tool_raw!r} -> row skipped")
            continue
        where = f"{SHEET}!row {n}"
        region = normalize_region(r[0] if len(r) > 0 else None, ctx, where)
        tool_type = cell_text(r[2] if len(r) > 2 else None).upper()
        if tool_code in seen_tools:
            ctx.warn(f"{SHEET}!row {n}: tool {tool_code} already listed at row {seen_tools[tool_code]} "
                     f"(this row: region {region!r}, type {tool_type!r}); region/type taken from the first row, cells merged")
        else:
            seen_tools[tool_code] = n
            ctx.see_tool("contact_chip", tool_code, raw=cell_text(tool_raw), region=region or None, tool_type=tool_type or None)
        for i, d in dates.items():
            if i >= len(r) or blank(r[i]):
                continue
            v, note = map_value(r[i], ctx)
            values[v or f"note:{note}"] += 1
            rows.append([tool_code, d, v, note])
    # duplicates (same tool listed twice with values on the same date): keep first, report
    dedup = {}
    for row in rows:
        key = (row[0], row[1])
        if key in dedup:
            ctx.warn(f"{SHEET}: duplicate value for {row[0]} on {row[1]} ({dedup[key][2] or dedup[key][3]!r} vs {row[2] or row[3]!r}); kept first")
            continue
        dedup[key] = row
    rows = sorted(dedup.values(), key=lambda x: (x[0], x[1]))
    out = Path(out_dir) / "contact_chip.csv"
    n = write_csv(out, HEADER, rows)
    ctx.count("contact_chip.csv", n)
    print(f"  contact_chip.csv: {n} rows, {len(seen_tools)} tools, {len(dates)} date columns "
          f"({min(dates.values()) if dates else '-'} .. {max(dates.values()) if dates else '-'}), values {dict(values)}")
    return n


def main(argv=None):
    ap = add_common_args(argparse.ArgumentParser(description=__doc__))
    a = ap.parse_args(argv)
    ctx = Context()
    run(a.src, a.out, ctx)
    ctx.save("contact_chip")
    print("  unmapped:", {k: dict(v) for k, v in ctx.unmapped.items() if v})


if __name__ == "__main__":
    main()
