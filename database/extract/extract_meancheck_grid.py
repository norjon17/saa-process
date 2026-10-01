"""Extract the 'Rev2.1' sheet of the First Mean Check Monitoring workbook
-> seed/meancheck_daily.csv and seed/meancheck_daily_excel_overall.csv

The sheet is huge (5,516 columns) so it is streamed straight from the xlsx zip
(xl/worksheets/sheetN.xml) with ElementTree.iterparse, exactly like
_ref_load_rev21.py.  Only the rows we need are kept in memory:
    row 12          dates (Excel serials), one per 8-column block from column E
    rows 26..76     tool rows: B area, C tool, D tool type, then 8 cells per day
    rows 81..128    Excel's computed Overall per (tool, day) - verification file
"""
from __future__ import annotations

import argparse
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    Context, MEANCHECK_XLSX, add_common_args, cell_text, normalize_daily_cell, normalize_region, normalize_tool,
    src_path, write_csv,
)

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
RNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
SHEET = "Rev2.1"
DATE_ROW = 12
TOOL_ROWS = range(26, 77)
OVERALL_ROWS = range(81, 129)
FIRST_BLOCK_COL = 5   # column E
BLOCK = 8
CELLS = ["vs_sat_psa_l", "vs_sat_psa_r", "vs_sat_rsa_l", "vs_sat_rsa_r",
         "vs_tgt_psa_l", "vs_tgt_psa_r", "vs_tgt_rsa_l", "vs_tgt_rsa_r"]
HEADER = ["tool_code", "date"] + CELLS + ["source_ref"]
OVERALL_HEADER = ["tool_code", "date", "excel_overall"]
OVERALL_VALUES = {"passed": "Passed", "failed": "Failed", "idle": "IDLE", "mb": "MB"}


def col2n(s: str) -> int:
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def sheet_part(z: zipfile.ZipFile, name: str) -> str:
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
    for sh in wb.iter(NS + "sheet"):
        if sh.get("name") == name:
            tgt = rels[sh.get(RNS + "id")]
            return tgt if tgt.startswith("xl/") else "xl/" + tgt.lstrip("/")
    raise KeyError(f"sheet {name!r} not found in {z.filename}")


def load_shared_strings(z: zipfile.ZipFile) -> list[str]:
    ss: list[str] = []
    if "xl/sharedStrings.xml" not in z.namelist():
        return ss
    for _ev, el in ET.iterparse(z.open("xl/sharedStrings.xml")):
        if el.tag == NS + "si":
            ss.append("".join(t.text or "" for t in el.iter(NS + "t")))
            el.clear()
    return ss


def stream_cells(z: zipfile.ZipFile, part: str, ss: list[str], want_rows: set[int]):
    """Yield (row, col, value_text) for valued cells in want_rows."""
    cell_re = re.compile(r"([A-Z]+)(\d+)")
    for _ev, el in ET.iterparse(z.open(part)):
        if el.tag == NS + "c":
            m = cell_re.match(el.get("r"))
            row = int(m.group(2))
            if row in want_rows:
                v = el.find(NS + "v")
                isv = el.find(NS + "is")
                if v is not None or isv is not None:
                    t = el.get("t")
                    if v is not None:
                        val = v.text
                        if t == "s" and val is not None:
                            val = ss[int(val)]
                    else:
                        val = "".join(x.text or "" for x in isv.iter(NS + "t"))
                    yield row, col2n(m.group(1)), val
            el.clear()
        elif el.tag == NS + "row":
            el.clear()


def run(src_dir, out_dir, ctx: Context) -> tuple[int, int]:
    path = src_path(src_dir, MEANCHECK_XLSX)
    z = zipfile.ZipFile(path)
    part = sheet_part(z, SHEET)
    ss = load_shared_strings(z)
    want = {DATE_ROW} | set(TOOL_ROWS) | set(OVERALL_ROWS)
    grid: dict[int, dict[int, str]] = {}
    for row, col, val in stream_cells(z, part, ss, want):
        if val is None or val == "":
            continue
        grid.setdefault(row, {})[col] = val
    print(f"  {SHEET} ({part}): kept {sum(len(v) for v in grid.values())} valued cells in {len(grid)} rows")

    # ---- dates
    ctx.decide("Rev2.1: row-12 date cells are trusted as-is (the sequence jumps twice); a (tool, date) that occurs in two blocks "
               "keeps the block with more filled cells; combined tool cells like 'SAA78H/SAA68H' are skipped")
    dates: dict[int, str] = {}   # block start col -> iso date
    prev = None
    for col in sorted(grid.get(DATE_ROW, {})):
        if col < FIRST_BLOCK_COL:
            continue
        raw = grid[DATE_ROW][col]
        if (col - FIRST_BLOCK_COL) % BLOCK != 0:
            ctx.warn(f"{SHEET}!{get_column_letter(col)}{DATE_ROW}: date cell not on an 8-column block boundary -> ignored")
            continue
        try:
            d = from_excel(float(raw)).date()
        except (TypeError, ValueError):
            ctx.warn(f"{SHEET}!{get_column_letter(col)}{DATE_ROW}: date cell {raw!r} is not an Excel serial -> block ignored")
            continue
        if prev is not None and (d - prev[1]).days != 1:
            ctx.warn(f"{SHEET}!{get_column_letter(col)}{DATE_ROW}: date out of sequence: {prev[1].isoformat()} "
                     f"(col {get_column_letter(prev[0])}) -> {d.isoformat()}; trusting the cell value")
        prev = (col, d)
        dates[col] = d.isoformat()
    print(f"  {len(dates)} day blocks: {min(dates.values())} .. {max(dates.values())}")

    # ---- tool rows
    rows: dict[tuple[str, str], tuple[list, int]] = {}
    cell_counts: Counter = Counter()
    tools_seen: dict[str, int] = {}
    dup_stats: dict[tuple[str, int, int], Counter] = defaultdict(Counter)
    for r in TOOL_ROWS:
        cells = grid.get(r, {})
        tool_raw = cell_text(cells.get(3))
        if not tool_raw:
            if any(c >= FIRST_BLOCK_COL for c in cells):
                ctx.warn(f"{SHEET}!row {r}: no tool code in column C but {sum(1 for c in cells if c >= FIRST_BLOCK_COL)} data cells -> row skipped")
            continue
        if "/" in tool_raw or "\\" in tool_raw:
            ctx.warn(f"{SHEET}!C{r}: combined tool cell {tool_raw!r} -> row skipped")
            continue
        tool_code = normalize_tool(tool_raw, ctx)
        if tool_code is None:
            ctx.warn(f"{SHEET}!C{r}: unresolved tool {tool_raw!r} -> row skipped")
            continue
        region = normalize_region(cells.get(2), ctx, f"{SHEET}!B{r}")
        tool_type = cell_text(cells.get(4)).upper()
        if tool_code in tools_seen:
            ctx.warn(f"{SHEET}!C{r}: tool {tool_code} also at row {tools_seen[tool_code]}; on duplicate (tool, date) the block with more filled cells wins")
        else:
            tools_seen[tool_code] = r
            ctx.see_tool("rev21", tool_code, raw=tool_raw, region=region or None, tool_type=tool_type or None)
        for col, date in dates.items():
            vals = []
            for k in range(BLOCK):
                v = normalize_daily_cell(cells.get(col + k), ctx)
                vals.append(v)
            filled = sum(1 for v in vals if v)
            if not filled:
                continue
            for v in vals:
                if v:
                    cell_counts[v] += 1
            ref = f"rev21:{r}:{get_column_letter(col)}"
            key = (tool_code, date)
            if key in rows:
                old_vals, old_filled = rows[key]
                old_row = int(old_vals[-1].split(":")[1])
                if filled > old_filled:
                    dup_stats[(tool_code, old_row, r)]["new"] += 1
                    rows[key] = ([tool_code, date] + vals + [ref], filled)
                else:
                    dup_stats[(tool_code, old_row, r)]["first"] += 1
                continue
            rows[key] = ([tool_code, date] + vals + [ref], filled)
    for (tool_code, r_old, r_new), st in sorted(dup_stats.items()):
        ctx.warn(f"{SHEET}: {tool_code} listed at rows {r_old} and {r_new}: {st['first'] + st['new']} duplicate (tool, date) blocks - "
                 f"kept the block with more filled cells (row {r_old} won {st['first']}x, row {r_new} won {st['new']}x)")
    out_rows = sorted((v[0] for v in rows.values()), key=lambda x: (x[0], x[1]))
    out = Path(out_dir) / "meancheck_daily.csv"
    n = write_csv(out, HEADER, out_rows)
    ctx.count("meancheck_daily.csv", n)
    ds = [r[1] for r in out_rows]
    print(f"  meancheck_daily.csv: {n} rows, {len(tools_seen)} tools, dates {min(ds)} .. {max(ds)}")
    print(f"  cell counts: {dict(cell_counts.most_common())}")

    # ---- Excel's Overall rows (verification)
    orows: dict[tuple[str, str], list] = {}
    ovals: Counter = Counter()
    for r in OVERALL_ROWS:
        cells = grid.get(r, {})
        tool_raw = cell_text(cells.get(3))
        if not tool_raw:
            continue
        if "/" in tool_raw or "\\" in tool_raw:
            ctx.warn(f"{SHEET}!C{r} (overall): combined tool cell {tool_raw!r} -> row skipped")
            continue
        tool_code = normalize_tool(tool_raw, ctx)
        if tool_code is None:
            ctx.warn(f"{SHEET}!C{r} (overall): unresolved tool {tool_raw!r} -> row skipped")
            continue
        for col, date in dates.items():
            raw = cell_text(cells.get(col))
            if not raw:
                continue
            v = OVERALL_VALUES.get(raw.casefold(), raw)
            ovals[v] += 1
            key = (tool_code, date)
            if key in orows:
                if orows[key][2] != v:
                    ctx.warn(f"{SHEET} overall: duplicate ({tool_code}, {date}) {orows[key][2]!r} vs {v!r} (row {r}); kept first")
                continue
            orows[key] = [tool_code, date, v]
    oout_rows = sorted(orows.values(), key=lambda x: (x[0], x[1]))
    oout = Path(out_dir) / "meancheck_daily_excel_overall.csv"
    on = write_csv(oout, OVERALL_HEADER, oout_rows)
    ctx.count("meancheck_daily_excel_overall.csv", on)
    print(f"  meancheck_daily_excel_overall.csv: {on} rows, values {dict(ovals.most_common())}")
    return n, on


def main(argv=None):
    ap = add_common_args(argparse.ArgumentParser(description=__doc__))
    a = ap.parse_args(argv)
    ctx = Context()
    run(a.src, a.out, ctx)
    ctx.save("meancheck_grid")
    print("  unmapped:", {k: dict(v) for k, v in ctx.unmapped.items() if v})


if __name__ == "__main__":
    main()
